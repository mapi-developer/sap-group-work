"""Build reports/final_report.docx from reports/final_report.md.

Supports the Markdown subset used in the report: headings, paragraphs with **bold** / *italic* /
`code`, tables, bullet and numbered lists, code blocks, images and the mermaid workflow (drawn as a PNG).

    python src/build_report.py
"""
import re
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from sap_data import FIG, ROOT

MD = ROOT / "reports" / "final_report.md"
OUT = ROOT / "reports" / "final_report.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
GREY = RGBColor(0x52, 0x51, 0x4E)
TEXT_WIDTH = Inches(6.3)
MAX_IMG_HEIGHT = Inches(4.3)


# --------------------------------------------------------------------------- workflow diagram
def draw_workflow(path):
    steps = ["SAP S/4HANA\nSE16N", "Excel exports\nEKKO · EKPO\nLFA1 · EKBE", "Parquet cache\nneeded columns",
             "EDA\nper table", "Join + US scope\n884 PO items", "Q1 … Q8\nanalysis", "Final\nreport"]
    fig, ax = plt.subplots(figsize=(14, 1.7))
    w, gap = 1.8, 0.35
    for i, s in enumerate(steps):
        x = i * (w + gap)
        ax.add_patch(FancyBboxPatch((x, 0), w, 1, boxstyle="round,pad=0.02,rounding_size=0.12",
                                    fc="#e8f0fb", ec="#2a78d6", lw=1.5))
        ax.text(x + w / 2, 0.5, s, ha="center", va="center", fontsize=10, color="#1f3a5f")
        if i:
            ax.add_patch(FancyArrowPatch((x - gap + 0.03, 0.5), (x - 0.03, 0.5), arrowstyle="-|>",
                                         mutation_scale=14, color="#52514e", lw=1.4))
    ax.set_xlim(-0.05, len(steps) * (w + gap) - gap + 0.05)
    ax.set_ylim(-0.08, 1.08)
    ax.axis("off")
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# --------------------------------------------------------------------------- low-level helpers
# OOXML requires child elements in schema order, so new ones are inserted at their slot.
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
               "tblLook", "tblCaption", "tblDescription"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]


def insert_ordered(parent, el, order):
    rank = order.index(el.tag.split("}")[1])
    for child in parent:
        name = child.tag.split("}")[1]
        if name in order and order.index(name) > rank:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def shade(cell, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    insert_ordered(cell._tc.get_or_add_tcPr(), shd, TCPR_ORDER)


def para_border(p, side, color="1F3A5F", size=8, space=4):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = insert_ordered(pPr, OxmlElement("w:pBdr"), PPR_ORDER)
    el = OxmlElement(f"w:{side}")
    for k, v in {"val": "single", "sz": str(size), "space": str(space), "color": color}.items():
        el.set(qn(f"w:{k}"), v)
    bdr.append(el)


def para_shade(p, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    insert_ordered(p._p.get_or_add_pPr(), shd, PPR_ORDER)


def table_borders(table, color="D0D5DD"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{side}")
        for k, v in {"val": "single", "sz": "4", "space": "0", "color": color}.items():
            el.set(qn(f"w:{k}"), v)
        borders.append(el)
    insert_ordered(tblPr, borders, TBLPR_ORDER)


def add_field(run, instr):
    for tag, text in [("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*\s][^*]*?\*(?!\*))")


def add_inline(p, text, size=None, color=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"
            r.font.size = Pt((size or 10.5) - 1)
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            continue
        elif part.startswith("*"):
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color


# --------------------------------------------------------------------------- document setup
def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

    st = doc.styles
    normal = st["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in [("Heading 1", 18, NAVY, 6, 10), ("Heading 2", 13.5, NAVY, 14, 6),
                                             ("Heading 3", 12, BLUE, 14, 6), ("Title", 30, NAVY, 0, 6)]:
        s = st[name]
        s.font.name = "Calibri"
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            s.element.rPr.rFonts.attrib.pop(qn(attr), None)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        st[name].paragraph_format.space_after = Pt(3)


def header_footer(doc):
    sec = doc.sections[-1]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_inline(hp, "Global Bike · US procurement analysis 2026", size=8.5, color=GREY)
    para_border(hp, "bottom", color="D0D5DD", size=4, space=2)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page "); r.font.size = Pt(8.5); r.font.color.rgb = GREY
    r = fp.add_run(); r.font.size = Pt(8.5); r.font.color.rgb = GREY
    add_field(r, "PAGE")


# --------------------------------------------------------------------------- block renderers
def add_table(doc, rows, aligns):
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_borders(t)
    # column widths proportional to content length, with a floor
    lens = [max(min(len(re.sub(r"[*`]", "", r[c])), 60) for r in rows) + 4 for c in range(ncol)]
    widths = [TEXT_WIDTH * l / sum(lens) for l in lens]
    for i, row in enumerate(rows):
        for c, txt in enumerate(row):
            cell = t.cell(i, c)
            cell.width = widths[c]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if aligns[c] == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i == 0:
                shade(cell, "1F3A5F")
                add_inline(p, txt, size=9, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
            else:
                if i % 2 == 0:
                    shade(cell, "F3F6FA")
                add_inline(p, txt, size=9)
    for row in t.rows:  # header repeats, rows don't split across pages
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit"); trPr.append(el)
    hdr = OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(hdr)
    # cell padding
    mar = OxmlElement("w:tblCellMar")
    for side, v in [("top", 50), ("left", 90), ("bottom", 50), ("right", 90)]:
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); mar.append(el)
    insert_ordered(t._tbl.tblPr, mar, TBLPR_ORDER)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc, lines):
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
        pf.left_indent = Cm(0.3); pf.keep_with_next = i < len(lines) - 1
        para_shade(p, "F4F4F2")
        para_border(p, "left", color="2A78D6", size=12, space=6)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"; r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


FIG_NO = [0]


def restart_numbering(doc, p):
    """Give a numbered list its own w:num so it starts at 1 again."""
    numbering = doc.part.numbering_part.element
    style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    override = OxmlElement("w:lvlOverride"); override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride"); start.set(qn("w:val"), "1")
    override.append(start); num.append(override)
    return num.numId


def set_num(p, num_id):
    numPr = p._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = num_id


def add_image(doc, path, caption):
    with Image.open(path) as im:
        w, h = im.size
    width = TEXT_WIDTH
    if width * h / w > MAX_IMG_HEIGHT:
        width = int(MAX_IMG_HEIGHT * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=width)
    FIG_NO[0] += 1
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    add_inline(c, f"Figure {FIG_NO[0]}. {caption}", size=8.5, color=GREY)
    c.runs[0].italic = True


def title_page(doc, title, subtitle):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.add_run(title.replace(": final report", ""))
    p = doc.add_paragraph()
    add_inline(p, "Final report", size=16, color=BLUE, bold=True)
    para_border(p, "bottom", color="2A78D6", size=12, space=8)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    add_inline(p, subtitle.strip("*"), size=11.5, color=GREY)
    p = doc.add_paragraph()
    add_inline(p, "US purchasing organization US00 · Reporting period Oct 2025 – Sep 2026 · 22 October 2026",
               size=10, color=GREY)
    for _ in range(3):
        doc.add_paragraph()


# --------------------------------------------------------------------------- markdown parser
def parse_table(lines):
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines]
    sep = rows.pop(1)
    aligns = ["right" if s.endswith(":") and not s.startswith(":") else "left" for s in sep]
    return rows, aligns


def build():
    workflow = FIG / "workflow.png"
    draw_workflow(workflow)

    doc = Document()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    setup(doc)
    header_footer(doc)

    lines = MD.read_text(encoding="utf-8").splitlines()
    i, n = 0, len(lines)
    first_h2 = True
    num_id = None
    title = subtitle = None
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("# "):
            title = s[2:]
            subtitle = lines[i + 2].strip() if lines[i + 2].strip().startswith("*") else ""
            title_page(doc, title, subtitle)
            i += 3 if subtitle else 1
            continue
        if s.startswith("```"):
            lang = s[3:].strip()
            j = i + 1
            while not lines[j].strip().startswith("```"):
                j += 1
            block = lines[i + 1:j]
            if lang == "mermaid":
                add_image(doc, workflow, "Analysis workflow from SAP extraction to the report")
            else:
                add_code(doc, block)
            i = j + 1
            continue
        if s.startswith("|"):
            j = i
            while j < n and lines[j].strip().startswith("|"):
                j += 1
            add_table(doc, *parse_table(lines[i:j]))
            i = j
            continue
        if s.startswith("#### "):
            doc.add_heading(s[5:], level=3); i += 1; continue
        if s.startswith("### "):
            doc.add_heading(s[4:], level=2); i += 1; continue
        if s.startswith("## "):
            h = doc.add_heading(level=1)
            if not first_h2:
                h.paragraph_format.page_break_before = True
            first_h2 = False
            h.add_run(s[3:])
            para_border(h, "bottom", color="1F3A5F", size=8, space=4)
            i += 1
            continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            add_image(doc, MD.parent / m.group(2), m.group(1))
            i += 1
            continue
        lm = re.match(r"(- |\d+\. )(.*)", s)
        if lm:
            text, j = lm.group(2), i + 1
            while j < n and lines[j].startswith("  ") and lines[j].strip():
                text += " " + lines[j].strip(); j += 1
            p = doc.add_paragraph(style="List Bullet" if lm.group(1) == "- " else "List Number")
            if lm.group(1) != "- ":
                if lm.group(1) == "1. ":
                    num_id = restart_numbering(doc, p)
                set_num(p, num_id)
            add_inline(p, text)
            i = j
            continue
        # paragraph: join until blank / block start
        text, j = s, i + 1
        while j < n and lines[j].strip() and not re.match(r"(#|\||```|!\[|- |\d+\. |---)", lines[j].strip()):
            text += " " + lines[j].strip(); j += 1
        p = doc.add_paragraph()
        if text.startswith("*") and text.endswith("*") and not text.startswith("**"):
            add_inline(p, text, size=9, color=GREY)
        else:
            add_inline(p, text)
        i = j

    doc.core_properties.title = title or "Final report"
    doc.core_properties.subject = "Global Bike US procurement analysis 2026"
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
